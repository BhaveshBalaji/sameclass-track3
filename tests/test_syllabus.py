from io import BytesIO
from unittest.mock import Mock, MagicMock
import pytest
import requests
from class_search_api.api import ASUClassSearchClient
from class_search_api.syllabus import prior_terms, lookup_syllabus, safe_download_url, extract_document_text, download_document
from class_search_api.transport import SourceError
from conftest import section


URL = 'https://class-syllabus-prod.s3.dualstack.us-west-2.amazonaws.com/test.pdf'


def page(*rows, complete=True):
    return {'classes': list(rows), 'complete': complete}


def test_prior_terms_across_year_boundary():
    assert prior_terms('2271', 4) == ['2267', '2264', '2261', '2257']


def test_historical_same_professor_beats_current_different_professor():
    client = Mock()
    client.search_all_classes.side_effect = [page(section('1', professor='Other Person')),
                                            page(section('2', term='2267'))]
    client.get_syllabus_url.return_value = URL
    result = lookup_syllabus(client, subject='CSE', catalog_nbr='575', term='2271',
                             professor='Ada Lovelace', fallback_terms=['2267'], include_text=False)
    assert result['source_term'] == '2267' and result['same_professor']
    assert result['historical'] and 'Historical' in result['warning']
    client.get_syllabus_url.assert_called_once_with('2', '2267')


def test_different_professor_fallback_is_labeled():
    client = Mock()
    client.search_all_classes.side_effect = [page(), page(section('2', term='2267', professor='Other Person'))]
    client.get_syllabus_url.return_value = URL
    result = lookup_syllabus(client, subject='CSE', catalog_nbr='575', term='2271',
                             professor='Ada', fallback_terms=['2267'], include_text=False)
    assert result['same_professor'] is False and result['historical']
    assert 'Different instructor' in result['warning']


def test_lookup_can_return_not_found_incomplete_or_unavailable():
    client = Mock()
    client.search_all_classes.return_value = page()
    params = dict(subject='CSE', catalog_nbr='575', term='2271', fallback_terms=[], include_text=False)
    assert lookup_syllabus(client, **params)['status'] == 'not_found'
    client.search_all_classes.return_value = page(complete=False)
    assert lookup_syllabus(client, **params)['status'] == 'incomplete'
    client.search_all_classes.side_effect = SourceError('ASU', 'blocked')
    assert lookup_syllabus(client, **params)['status'] == 'unavailable'


def test_lookup_excludes_prefix_courses_and_preserves_missing_syllabus():
    client = Mock()
    client.search_all_classes.return_value = page(section(catalog='575A'))
    result = lookup_syllabus(client, subject='CSE', catalog_nbr='575', term='2271', fallback_terms=[], include_text=False)
    assert result['status'] == 'not_found'
    client.get_syllabus_url.assert_not_called()
    client.search_all_classes.return_value = page(section())
    client.get_syllabus_url.side_effect = SourceError('Syllabus', 'missing', status=404)
    assert lookup_syllabus(client, subject='CSE', catalog_nbr='575', term='2271', fallback_terms=[], include_text=False)['status'] == 'not_found'


@pytest.mark.parametrize('terms', [['2271'], ['2274'], ['bad']])
def test_fallback_cannot_be_current_future_or_invalid(terms):
    with pytest.raises(ValueError):
        lookup_syllabus(Mock(), subject='CSE', catalog_nbr='575', term='2271', fallback_terms=terms)


@pytest.mark.parametrize('url', ['http://asu.edu/a', 'https://127.0.0.1/a', 'https://evil.com/a',
    'https://asu.edu.evil.com/a', 'https://user:password@asu.edu/a', 'https://asu.edu:8000/a'])
def test_download_rejects_unsafe_hosts(url):
    with pytest.raises(SourceError):
        safe_download_url(url)


def test_real_404_is_not_reported_as_upstream_outage():
    session = Mock()
    response = requests.Response()
    response.status_code = 404
    session.request.return_value.raise_for_status.side_effect = requests.HTTPError(response=response)
    client = ASUClassSearchClient(session=session)
    with pytest.raises(SourceError) as exc:
        client.get_syllabus_url('1', '2271')
    assert exc.value.status == 404


def test_text_and_html_extraction_truncation():
    result = extract_document_text(b'<html><script>bad()</script><p>Assignments 30%</p></html>', 'text/html')
    assert 'Assignments 30%' in result['text'] and 'bad()' not in result['text']
    result = extract_document_text(b'abcdef', 'text/plain', 3)
    assert result['text'] == 'abc' and result['truncated']
    with pytest.raises(SourceError):
        extract_document_text(b'unknown binary', 'application/octet-stream')


def test_blank_pdf_needs_ocr_and_docx_tables_are_extracted():
    from pypdf import PdfWriter
    from docx import Document
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = BytesIO()
    writer.write(stream)
    result = extract_document_text(stream.getvalue())
    assert result['ocr_required'] and result['pages'][0]['page'] == 1
    doc = Document()
    doc.add_paragraph('Course schedule')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Exam', '40%'
    stream = BytesIO()
    doc.save(stream)
    assert 'Exam | 40%' in extract_document_text(stream.getvalue())['text']


def test_download_redirect_target_is_revalidated(monkeypatch):
    session = MagicMock()
    response = session.get.return_value.__enter__.return_value
    response.status_code = 302
    response.headers = {'Location': 'http://127.0.0.1/private'}
    monkeypatch.setattr(requests, 'Session', lambda: session)
    session.__enter__ = Mock(return_value=session)
    session.__exit__ = Mock(return_value=False)
    with pytest.raises(SourceError):
        download_document(Mock(timeout=10), URL)
    assert session.get.call_count == 1


def test_syllabus_routes_redirect_json_and_missing(asu_http):
    http, client = asu_http
    client.get_syllabus_url.return_value = URL
    result = http.get('/api/syllabus_url?term=2271&class_nbr=123', follow_redirects=False)
    assert result.status_code == 307 and result.headers['location'] == URL
    result = http.get('/api/syllabus_url?term=2271&class_nbr=123&redirect=false')
    assert result.json()['download_url'] == URL
    client.get_syllabus_url.side_effect = SourceError('ASU Syllabus', 'missing', status=404)
    assert http.get('/api/syllabus_url?term=2271&class_nbr=123').status_code == 404
